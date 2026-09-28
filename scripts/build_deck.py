"""Build Vertex_Selector_Approach_Deck.pptx. pip install python-pptx; python scripts/build_deck.py [out.pptx]"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x1F, 0x6F, 0x8B)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x65, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xD5, 0xDB, 0xE0)
ROW = RGBColor(0xF4, 0xF7, 0xF8)
W, H = Inches(13.333), Inches(7.5)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Vertex_Selector_Approach_Deck.pptx"


def _set_run(run, text, size, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"


def add_text(slide, l, t, w, h, lines, default_size=18, default_color=INK):
    """lines: str or list of (text, size, color, bold)."""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    if isinstance(lines, str):
        lines = [(lines, default_size, default_color, False)]
    for i, item in enumerate(lines):
        if isinstance(item, str):
            item = (item, default_size, default_color, False)
        text, size, color, bold = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.add_run()
        _set_run(run, text, size, color, bold)
    return tb


def bar(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
    s.fill.solid()
    s.fill.fore_color.rgb = TEAL
    s.line.fill.background()
    foot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, H - Inches(0.32), W, Inches(0.32))
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY
    foot.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.5), H - Inches(0.30), Inches(12), Inches(0.28))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, "Vertex Smart Categorization  ·  Search Result Selector  ·  Confidential working session", 10, WHITE)


def title_block(slide, kicker, title, sub=None):
    bar(slide)
    add_text(slide, Inches(0.55), Inches(0.32), Inches(12), Inches(0.35),
             [(kicker.upper(), 12, TEAL, True)])
    add_text(slide, Inches(0.55), Inches(0.58), Inches(12.2), Inches(0.7),
             [(title, 28, NAVY, True)])
    if sub:
        add_text(slide, Inches(0.55), Inches(1.18), Inches(12.2), Inches(0.4),
                 [(sub, 14, MUTED, False)])


def bullets(slide, l, t, w, h, items, size=16):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        _set_run(run, "•  " + item, size, INK)
    return tb


def table(slide, l, t, w, h, rows, col_w=None):
    n_rows, n_cols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_rows, n_cols, l, t, w, h)
    tbl = shp.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = cw
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = p.add_run()
            _set_run(run, str(val), 13, WHITE if r == 0 else INK, bold=(r == 0 or c == 0))
            fill = NAVY if r == 0 else (ROW if r % 2 == 0 else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
    return shp


def new(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return slide


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 title
    s = new(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    add_text(s, Inches(0.7), Inches(1.8), Inches(12), Inches(0.4),
             [("STAKEHOLDER UPDATE", 14, TEAL, True)])
    add_text(s, Inches(0.7), Inches(2.2), Inches(12), Inches(1.4),
             [("Search Result Selector", 40, WHITE, True),
              ("Smart Categorization Engine", 22, RGBColor(0xC5, 0xD4, 0xDE), False)])
    add_text(s, Inches(0.7), Inches(5.6), Inches(12), Inches(0.8),
             [("Which provided search results are about the uploaded product —", 16, WHITE, False),
              ("and nothing else, until that decision is trustworthy.", 16, WHITE, False)])

    # 2 why
    s = new(prs)
    title_block(s, "The risk", "A wrong link is a wrong tax category",
                "This step sits between a messy ERP title and every downstream enrichment.")
    bullets(s, Inches(0.6), Inches(1.8), Inches(12), Inches(4.8), [
        "Clients upload title + description only. Tax departments already “Google it” to find ingredients.",
        "We will scrape the pages we trust. False pages → false carbonation / juice % / sweetener / alcohol.",
        "Those pages are also shown in the product UI. Customers will not forgive a sibling flavor or a listing page.",
        "100k+ SKU retailers. One bad heuristic, multiplied, is a support queue — and an O-Series miss.",
        "Grey area is real: size does not matter; flavor and formulation do; “Apple” should link to nothing.",
    ])

    # 3 scope
    s = new(prs)
    title_block(s, "Scope", "This case is the selection step only")
    table(s, Inches(0.55), Inches(1.7), Inches(12.2), Inches(2.4), [
        ["In this deliverable", "Not in this deliverable"],
        ["Ingest title, description, search_results", "Live OpenAI web search"],
        ["Emit llm_trusted_search_results (indices)", "Fetch / scrape the trusted pages"],
        ["Score against trusted_search_results", "Map an enriched description to a tax category"],
    ], col_w=[Inches(6.1), Inches(6.1)])
    add_text(s, Inches(0.55), Inches(4.5), Inches(12.2), Inches(2), [
        ("The brief attaches the search results. The hidden test file is the same shape.", 16, INK, False),
        ("A live search would return different URLs than the labeled indices we are scored on.", 16, INK, False),
        ("Web search belongs in the next stage of the enrichment pipeline, after this step is trusted.", 16, INK, False),
    ])

    # 4 approach
    s = new(prs)
    title_block(s, "Methodology", "One structured judgment per product, fail closed")
    bullets(s, Inches(0.6), Inches(1.75), Inches(12.1), Inches(5.2), [
        "One LLM call sees every result, so it can notice “3 of 10 are the exact flavor; the rest are siblings.”",
        "Verdicts, not a bare yes/no: match · different_variant · different_product · generic_or_listing.",
        "Same product = same brand + line + variant. Ignore size / pack. Honor flavor, diet, decaf, alcohol, form.",
        "Generic or ambiguous titles set product_is_specific=false and we return [].",
        "Guardrails: JSON schema with a closed enum of legal indices; injection flag; sanitize prompt tags; retries.",
        "Heuristic lexical baseline must be beaten — and gives us an offline path plus a disagreement sample.",
        "Every decision is written to *.audit.json (verdict, confidence, reason) for the code review and for tax.",
    ])

    # 5 metrics
    s = new(prs)
    title_block(s, "Metric choice", "Optimize F0.5 and clean-product rate")
    table(s, Inches(0.55), Inches(1.7), Inches(12.2), Inches(2.8), [
        ["Metric", "Why it is the one we show you"],
        ["F0.5", "Precision counts twice. A miss is coverage; a false link is a wrong category."],
        ["Clean-product rate", "Share of products with zero false links — the user-trust number."],
        ["Recall / coverage", "Reported as the cost of being strict, not the thing we maximize."],
        ["Correct abstain", "Did we stay silent on “Apple”, deli, bakery, unlinkable ERP junk?"],
    ], col_w=[Inches(3.2), Inches(9.0)])
    add_text(s, Inches(0.55), Inches(5.0), Inches(12.2), Inches(1.4),
             [("Ask of the business: do not make us chase recall if it drops clean-product rate.", 16, NAVY, True)])

    # 6 results
    s = new(prs)
    title_block(s, "Evidence", "train_50 — development numbers, not a launch claim",
                "Same 50 products informed the prompt. Hidden test is the generalization check.")
    table(s, Inches(0.55), Inches(1.75), Inches(12.2), Inches(2.0), [
        ["Backend", "Precision", "Recall", "F0.5", "Clean-product"],
        ["Heuristic (offline baseline)", "0.674", "0.800", "0.696", "0.400"],
        ["LLM  gpt-4.1-mini  conf ≥ 0.70", "0.796", "0.908", "0.816", "0.640"],
    ], col_w=[Inches(4.4), Inches(1.95), Inches(1.95), Inches(1.95), Inches(1.95)])
    bullets(s, Inches(0.6), Inches(4.15), Inches(12), Inches(2.4), [
        "LLM abstains more correctly on generic items (0.667 vs 0.167) and still finds ≥1 good link on 97% of labeled products.",
        "18 products still have at least one false link — the review set, not a reason to loosen the threshold tonight.",
        "Label ceiling is below 1.0: 12/50 products have an empty gold list; 7 of 82 duplicate-URL groups are labeled inconsistently. Some ‘errors’ will be label errors.",
        "Re-threshold from the audit file (`--conf`) without another paid call.",
    ], size=14)

    # 7 guardrails
    s = new(prs)
    title_block(s, "Production posture", "What we already refuse to do")
    bullets(s, Inches(0.6), Inches(1.75), Inches(12.1), Inches(5.2), [
        "Fail closed: API error, refusal, bad JSON, missing indices, unspecific product → trust nothing.",
        "The model cannot invent an index that was not in the input (schema enum).",
        "Snippet text that looks like “ignore previous instructions” is never trusted.",
        "One product cannot crash a batch. 8 workers, per-row isolation.",
        "Cache is keyed on model + prompt + schema so a meeting replay of train_50 is free and identical.",
        "No silent fallback to the heuristic in production. If OpenAI is down, we stop or return [].",
    ])

    # 8 tradeoffs
    s = new(prs)
    title_block(s, "Trade-offs", "What I am asking the business to accept")
    table(s, Inches(0.55), Inches(1.7), Inches(12.2), Inches(4.6), [
        ["Choice", "We gain", "We give up"],
        ["Precision > recall", "Fewer poisoned categories / UI links", "Some branded SKUs stay unenriched"],
        ["Snippet only (no fetch)", "Cost, latency, no extra dependency", "Nav-junk snippets we could have resolved"],
        ["Fail closed", "Never guess under an outage", "Empty enrichment when the model flakes"],
        ["LLM default", "Variant / listing judgment", "Key, cost, vendor dependency"],
        ["Ignore size; honor flavor", "Matches tax-component reality", "Disagrees with some human labelers"],
    ], col_w=[Inches(3.4), Inches(4.4), Inches(4.4)])

    # 9 requirements
    s = new(prs)
    title_block(s, "Requirements", "What must be true before this feeds O-Series")
    bullets(s, Inches(0.6), Inches(1.75), Inches(12.1), Inches(5.2), [
        "Pinned model + prompt version on every audit row. A prompt edit is a release.",
        "Confidence threshold tunable offline from the audit file.",
        "Human-review sample each release: false-link products and LLM-vs-lexical disagreements.",
        "Analyst override in the product (add/remove an index) logged as future gold.",
        "Proposed launch bar (open for debate): clean-product ≥ 0.80 and precision ≥ 0.85 on a held-out set.",
        "Confirmed cost / p95 latency envelope for a 100k-SKU job.",
        "Clear rule on whether customer SKU strings may be sent to OpenAI, and how long we keep audits.",
    ])

    # 10 roadmap
    s = new(prs)
    title_block(s, "Roadmap", "This step, then enrichment, then classification")
    table(s, Inches(0.55), Inches(1.7), Inches(12.2), Inches(4.4), [
        ["Horizon", "Work"],
        ["Now", "Selector + metrics + synthetics + guardrails + this review. Run the hidden test in this meeting."],
        ["2 weeks", "Second labeler on the 18 dirty train rows. Re-threshold. Page-fetch only if you want coverage back."],
        ["Next quarter", "Scrape trusted URLs → enriched description. Classify to a Vertex tax category. Online monitors."],
        ["Later", "Client abbreviation tables. Active learning from overrides. Jurisdiction policy if it changes “same product”."],
    ], col_w=[Inches(2.4), Inches(9.8)])

    # 11 questions
    s = new(prs)
    title_block(s, "Clarifying questions", "What I would ask if I owned this at Vertex")
    q = [
        "1. Whose labels are gold when annotators disagree? 7 of 82 duplicate-URL groups are inconsistent.",
        "2. Does the hidden file include trusted_search_results, or only predictions on screen?",
        "3. RTD vs mix vs bags vs concentrate of the same flavor — match or different product?",
        "4. Diet / decaf / lite / spiked: always a different variant, even if the title is silent?",
        "5. Produce / deli / bakery / “Apple”: abstain in the UI, or still show a generic page?",
        "6. Same SKU on Amazon + Kroger + manufacturer — trust all matches, or one canonical URL?",
        "7. Show every trusted index, top-N, or a three-way auto / suggest / abstain?",
        "8. May analysts override a link, and may we treat overrides as future training labels?",
        "9. p95 latency and budget per 100k SKUs? Batch API acceptable for bulk?",
        "10. Who owns the search step (query, engine, result count)? Can we change it?",
        "11. May titles (customer SKU strings) go to OpenAI, and how long do we retain audits?",
        "12. What clean-product / precision bar unlocks the rest of Smart Categorization?",
    ]
    bullets(s, Inches(0.55), Inches(1.65), Inches(12.2), Inches(5.4), q, size=14)

    # 12 live test
    s = new(prs)
    title_block(s, "Live test", "Drop the file in; we keep talking")
    add_text(s, Inches(0.6), Inches(1.75), Inches(12), Inches(1.4), [
        ("$env:OPENAI_API_KEY='sk-...'", 20, NAVY, True),
        ("python run.py test.json", 20, NAVY, True),
        ("python eval.py test.json          # same run, formatted headline numbers", 16, INK, False),
    ])
    bullets(s, Inches(0.6), Inches(3.6), Inches(12), Inches(2.8), [
        "Writes test.output.json + test.output.audit.json in the background.",
        "If labels are present: F0.5, precision, clean-product print immediately.",
        "If labels are omitted: predictions still write; metrics say scored: false. That is not a crash.",
        "No key: cached products replay; new products fail closed. Offline: --backend heuristic.",
    ], size=15)

    # 13 ask
    s = new(prs)
    title_block(s, "Ask", "What I need from this room")
    bullets(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(4.8), [
        "Agree that precision / clean-product is the launch metric — not raw recall.",
        "Answer the form / alcohol / abstain / top-N questions before we retune.",
        "Hand over the hidden test JSON so we can watch the contract hold on unseen data.",
        "Leave-behind: docs/product-owner-brief.md  ·  then a code walk of selector/.",
    ])

    dest = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else OUT
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(dest)
    print("wrote", dest)


if __name__ == "__main__":
    build()
